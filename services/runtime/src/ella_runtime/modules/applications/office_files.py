"""Create Word, Excel, and PowerPoint files inside the agent workspace."""

import hashlib
import os
from pathlib import Path
from typing import Any
from uuid import uuid4

from ella_runtime.modules.agent.contracts import ToolEvidence
from ella_runtime.storage_paths import default_data_dir


class OfficeUnavailable(RuntimeError):
    """Optional Office file libraries have not been installed."""


class OfficeFileTool:
    name: str
    description: str
    extension: str

    def __init__(self, root: Path | None = None) -> None:
        self.root = (root or default_data_dir() / "workspace").resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def _target(self, arguments: dict[str, Any]) -> tuple[Path, bool]:
        name = arguments.get("path")
        overwrite = arguments.get("overwrite", False)
        if not isinstance(name, str) or not name.strip() or Path(name).is_absolute():
            raise ValueError("文件路径必须是工作区内的相对路径")
        if not isinstance(overwrite, bool):
            raise TypeError("覆盖设置必须是布尔值")
        target = (self.root / name).resolve()
        if target == self.root or not target.is_relative_to(self.root):
            raise ValueError("文件路径超出工作区")
        if target.suffix.lower() != self.extension:
            raise ValueError(f"文件扩展名必须为 {self.extension}")
        return target, overwrite

    async def execute(self, arguments: dict[str, Any], *, action_id: str) -> ToolEvidence:
        target, overwrite = self._target(arguments)
        self._validate(arguments)
        if target.exists() and not overwrite:
            prior = await self.reconcile(arguments, action_id=action_id)
            if prior is not None and prior.verified:
                return prior
            raise FileExistsError("目标文件已存在，计划未允许覆盖")
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.parent / f".{target.stem}.{uuid4().hex}{self.extension}"
        try:
            self._create(temporary, arguments)
            checked = self._verify(temporary, arguments)
            if not checked:
                raise RuntimeError("Office 文件写入后内容核验失败")
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
        result = await self.reconcile(arguments, action_id=action_id)
        assert result is not None
        return result

    async def reconcile(self, arguments: dict[str, Any], *, action_id: str) -> ToolEvidence | None:
        target, _ = self._target(arguments)
        self._validate(arguments)
        if not target.is_file():
            return None
        verified = self._verify(target, arguments)
        return ToolEvidence(
            verified=verified,
            details={
                "path": str(target),
                "bytes": target.stat().st_size,
                "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
                "content_verified": verified,
            },
        )

    def _validate(self, arguments: dict[str, Any]) -> None:
        raise NotImplementedError

    def _create(self, path: Path, arguments: dict[str, Any]) -> None:
        raise NotImplementedError

    def _verify(self, path: Path, arguments: dict[str, Any]) -> bool:
        raise NotImplementedError


class SpreadsheetTool(OfficeFileTool):
    name = "office.create_spreadsheet"
    extension = ".xlsx"
    description = (
        "创建 Excel .xlsx 文件并重新打开核验。参数："
        '{"path":"表格.xlsx","sheet":"工作表名","rows":[["列一","列二"],["A",1]],'
        '"overwrite":false}。最多 200 行、20 列。'
    )

    def _validate(self, arguments: dict[str, Any]) -> None:
        sheet = arguments.get("sheet", "Sheet1")
        rows = arguments.get("rows")
        if (
            not isinstance(sheet, str)
            or not 1 <= len(sheet) <= 31
            or any(char in sheet for char in "[]:*?/\\")
        ):
            raise ValueError("工作表名称无效")
        if not isinstance(rows, list) or not 1 <= len(rows) <= 200:
            raise ValueError("表格行数必须在 1 到 200 之间")
        for row in rows:
            if not isinstance(row, list) or not 1 <= len(row) <= 20:
                raise ValueError("每行必须包含 1 到 20 列")
            if any(
                value is not None and type(value) not in {str, int, float, bool} for value in row
            ):
                raise TypeError("单元格只支持文本、数字、布尔值或空值")

    def _create(self, path: Path, arguments: dict[str, Any]) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise OfficeUnavailable("请安装 runtime[office] 依赖") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = arguments.get("sheet", "Sheet1")
        for row in arguments["rows"]:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4369AF")
        for column in sheet.columns:
            letter = column[0].column_letter
            widest = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[letter].width = min(max(widest + 2, 12), 50)
        workbook.save(path)

    def _verify(self, path: Path, arguments: dict[str, Any]) -> bool:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise OfficeUnavailable("请安装 runtime[office] 依赖") from exc
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet_name = arguments.get("sheet", "Sheet1")
            if workbook.sheetnames != [sheet_name]:
                return False
            actual = list(workbook[sheet_name].values)
            expected = [tuple(row) for row in arguments["rows"]]
            width = max(len(row) for row in expected)
            return [tuple(row[:width]) for row in actual] == [
                tuple(row) + (None,) * (width - len(row)) for row in expected
            ]
        finally:
            workbook.close()


class DocumentTool(OfficeFileTool):
    name = "office.create_document"
    extension = ".docx"
    description = (
        "创建 Word .docx 文件并重新打开核验。参数："
        '{"path":"文档.docx","title":"标题","paragraphs":["第一段","第二段"],'
        '"overwrite":false}。最多 100 段。'
    )

    def _validate(self, arguments: dict[str, Any]) -> None:
        title = arguments.get("title")
        paragraphs = arguments.get("paragraphs")
        if not isinstance(title, str) or not title.strip() or len(title) > 200:
            raise ValueError("文档标题无效")
        if (
            not isinstance(paragraphs, list)
            or len(paragraphs) > 100
            or any(not isinstance(text, str) or len(text) > 5000 for text in paragraphs)
        ):
            raise ValueError("文档段落无效")

    def _create(self, path: Path, arguments: dict[str, Any]) -> None:
        try:
            from docx import Document
        except ImportError as exc:
            raise OfficeUnavailable("请安装 runtime[office] 依赖") from exc
        document = Document()
        document.add_heading(arguments["title"], level=0)
        for text in arguments["paragraphs"]:
            document.add_paragraph(text)
        document.save(path)

    def _verify(self, path: Path, arguments: dict[str, Any]) -> bool:
        try:
            from docx import Document
        except ImportError as exc:
            raise OfficeUnavailable("请安装 runtime[office] 依赖") from exc
        document = Document(path)
        return [paragraph.text for paragraph in document.paragraphs] == [
            arguments["title"],
            *arguments["paragraphs"],
        ]


class PresentationTool(OfficeFileTool):
    name = "office.create_presentation"
    extension = ".pptx"
    description = (
        "创建 PowerPoint .pptx 文件并重新打开核验。参数："
        '{"path":"演示.pptx","slides":[{"title":"标题","bullets":["要点一","要点二"]}],'
        '"overwrite":false}。最多 20 页。'
    )

    def _validate(self, arguments: dict[str, Any]) -> None:
        slides = arguments.get("slides")
        if not isinstance(slides, list) or not 1 <= len(slides) <= 20:
            raise ValueError("幻灯片数量必须在 1 到 20 之间")
        for slide in slides:
            if (
                not isinstance(slide, dict)
                or not isinstance(slide.get("title"), str)
                or not slide["title"].strip()
                or len(slide["title"]) > 200
            ):
                raise TypeError("幻灯片标题无效")
            bullets = slide.get("bullets", [])
            if (
                not isinstance(bullets, list)
                or len(bullets) > 20
                or any(not isinstance(item, str) or len(item) > 1000 for item in bullets)
            ):
                raise ValueError("幻灯片要点无效")

    def _create(self, path: Path, arguments: dict[str, Any]) -> None:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise OfficeUnavailable("请安装 runtime[office] 依赖") from exc
        presentation = Presentation()
        for item in arguments["slides"]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = item["title"]
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for index, bullet in enumerate(item.get("bullets", [])):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = bullet
                paragraph.level = 0
        presentation.save(path)

    def _verify(self, path: Path, arguments: dict[str, Any]) -> bool:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise OfficeUnavailable("请安装 runtime[office] 依赖") from exc
        presentation = Presentation(path)
        if len(presentation.slides) != len(arguments["slides"]):
            return False
        for slide, expected in zip(presentation.slides, arguments["slides"], strict=True):
            if slide.shapes.title.text != expected["title"]:
                return False
            actual_bullets = [
                paragraph.text for paragraph in slide.placeholders[1].text_frame.paragraphs
            ]
            if actual_bullets == [""] and not expected.get("bullets", []):
                continue
            if actual_bullets != expected.get("bullets", []):
                return False
        return True
